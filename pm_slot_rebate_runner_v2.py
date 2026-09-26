# -*- coding: utf-8 -*-
"""Patron Management Slot Rebate test automation - v2 (sequential, gated steps).

Same workflow as the v1.x runner. Reads Excel-for-auto.xlsx, sheet REBATE,
keeps rows where SLOT REBATE 5% != 0, then for each row, strictly in order:

   1  pre-check: PM idle (no popup, main window enabled, logged in)
   2  open Find a Player (Ribbon "Find Player" button, fallback Ctrl+F)
   3  enter the Player ID (the field may still hold the previous ID)
   4  OK -> Find a Player closes
   5  wait until the main title shows this Player ID (profile loaded)
   6  handle profile popups until PM is quiet:
        System Messages -> log every line (incl. Player Stop Codes) -> Close
        Player Comment  -> Next until Close is enabled -> Close
   7  Options... -> Options menu opens
   8  Redeem Coupon... -> Coupon Redemption opens
   9  Competitor Coupon -> Competitor list and Amount become enabled
  10  select DAILY REBATE (5%) - 1
  11  enter the Slot Rebate amount from Excel
  12  OK or Cancel -> Coupon Redemption closes

At start the runner asks whether it may click OK (real redemption). Typing YES
clicks OK in step 12; any other answer clicks Cancel (test run, as before).
Players redeemed today (pm_redeemed_ledger.csv) are skipped in OK mode so a
rerun never redeems the same player twice. Steps are separated by a
STEP_DELAY_SECONDS pause.

Every step checks the PM state before acting (which popups are open, main
window enabled or blocked, correct player) and waits for the expected state
afterwards. Popups are found through Win32 (PM process ID + HWND); UIA is only
used where PM has no native control (Ribbon button, menu items, list items).

On the first error the runner stops immediately and leaves PM exactly as it
is, so a person can look at the screen. After the last Excel row has been
processed, the profile tabs opened by the runner are closed.

Timings in comments are the values measured with pm_recorder_v5.
"""

import csv
import ctypes
import re
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

try:
    import win32con
    import win32gui
    import win32process
    from pywinauto import Desktop, handleprops, mouse
    from pywinauto.controls.hwndwrapper import HwndWrapper
    from pywinauto.keyboard import send_keys
except ImportError:  # lets the step logic be imported and tested off Windows
    win32con = win32gui = win32process = None

RUNNER_VERSION = "2.0-sequential-gated"
BASE_DIR = Path(__file__).resolve().parent
EXCEL_FILE = BASE_DIR / "Excel-for-auto.xlsx"
LOG_FILE = BASE_DIR / "pm_automation_log.csv"
TEXT_LOG_FILE = BASE_DIR / "pm_runner_log.txt"
SHEET_NAME = "REBATE"
PLAYER_ID_COLUMN = "Player ID"
SLOT_REBATE_COLUMN = "SLOT REBATE 5%"
COMPETITOR_NAME = "DAILY REBATE (5%) - 1"
STEP_DELAY_SECONDS = 2.0   # pause before every step
LEDGER_FILE = BASE_DIR / "pm_redeemed_ledger.csv"
CLOSE_TABS_AT_END = True
FIND_SHORTCUT = "^f"
STOP_CODES_ITEM = "Player Stop Codes"

# Timeouts (seconds). Measured values from the v5 recordings in comments.
T_LOGIN = 90
T_FIND_OPEN = 3            # 0.21-0.31s
T_FIND_CLOSE = 5           # 0.02-0.16s
T_PROFILE_LOAD = 30        # 4.1-5.2s from OK to the new title
T_PROFILE_POPUPS = 60      # System Messages + paging through all comments
T_POPUP_CLOSE = 5          # 0.12-0.15s
T_AFTER_OK = 20            # not recorded yet: OK has never been clicked in a recording
T_COMMENT_PAGE = 3         # 0.11-0.2s per Next
T_MENU_OPEN = 3            # 0.15-0.18s
T_COUPON_OPEN = 5          # 0.21-0.22s
T_FIELDS_ENABLE = 3        # 0.13-0.24s after Competitor Coupon
T_VERIFY = 3
T_TAB_CLOSE = 5
QUIET_SECONDS = 1.0        # measured gap between System Messages and Player Comment: 0.04-0.05s
SAME_PLAYER_WAIT_SECONDS = 8.0      # profile load measured 4.1-5.2s
POLL = 0.1
MAX_COMMENT_PAGES = 200

# Native control IDs (identical in every recording).
ID_OK = 1
ID_CANCEL = 2
FIND_PLAYER_ID_EDIT = 1041
SYSMSG_CLOSE = 1
COMMENT_NEXT = 1002
COMMENT_PREVIOUS = 1003
COMMENT_CLOSE = 2
COMMENT_HEADER = 100
OPTIONS_BUTTON = 1074
COUPON_OUR_RADIO = 1481
COUPON_COMPETITOR_RADIO = 1484
COUPON_ID_EDIT = 1477
COUPON_COMPETITOR_COMBO = 1483
COUPON_AMOUNT_EDIT = 1107
COUPON_CLICKABLE_IDS = {COUPON_COMPETITOR_RADIO, COUPON_COMPETITOR_COMBO, COUPON_AMOUNT_EDIT, ID_CANCEL}
REDEEM_MENU_ITEM = "Redeem Coupon..."
FIND_RIBBON_BUTTON = "Find Player"

# Popup kinds.
LOGIN, FIND, SYSMSG, COMMENT, COUPON, MENU, DROPDOWN, UNKNOWN = (
    "LOGIN", "FIND", "SYSTEM_MESSAGES", "PLAYER_COMMENT", "COUPON", "MENU", "DROPDOWN", "UNKNOWN")
DIALOG_TITLES = {
    "find a player": FIND,
    "system messages": SYSMSG,
    "system message": SYSMSG,
    "player comment": COMMENT,
    "coupon redemption": COUPON,
}
LOGIN_TITLE_RE = re.compile(r"^Patron Management\s+(?:Log\s*on|Log\s*in)$", re.I)
MAIN_TITLE_RE = re.compile(r"^Patron Management(?: - .+)?$")
MAIN_CLASSES = {"XTPMainFrame"}
MIN_POPUP_SIZE = 9          # XTP menu shadows are 4px wide windows
IGNORED_CLASS_RE = re.compile(r"tooltip|shadow|PopupBubbleWnd|^IME$|MSCTFIME", re.I)
LOGIN_BUTTON_NAMES = {"login", "log in", "logon", "log on", "sign in", "ok"}
WM_MDIGETACTIVE = 0x0229


class StepError(RuntimeError):
    def __init__(self, step, message):
        super().__init__(message)
        self.step = step


def log(message):
    line = f"[{datetime.now().strftime('%H:%M:%S')}] {message}"
    print(line, flush=True)
    try:
        with TEXT_LOG_FILE.open("a", encoding="utf-8") as file:
            file.write(line + "\n")
    except OSError:
        pass


# ------------------------------------------------------------------ Excel

def original_excel_text(value):
    """Preserve Excel's value representation without padding or recalculation."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else format(value, "g")
    return str(value).strip()


def numeric_nonzero(value):
    if value is None or value == "":
        return False
    try:
        return float(value) != 0
    except (TypeError, ValueError):
        raise ValueError(f"Invalid rebate value: {value!r}")


def read_jobs():
    if not EXCEL_FILE.exists():
        raise FileNotFoundError(f"Excel file not found: {EXCEL_FILE}")
    wb = load_workbook(EXCEL_FILE, data_only=True, read_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise KeyError(f"Sheet not found: {SHEET_NAME}")
        rows = wb[SHEET_NAME].iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValueError("REBATE sheet is empty.")
        normalized = [str(x).strip().casefold() if x is not None else "" for x in headers]
        try:
            player_idx = normalized.index(PLAYER_ID_COLUMN.casefold())
            rebate_idx = normalized.index(SLOT_REBATE_COLUMN.casefold())
        except ValueError as exc:
            raise ValueError(
                f"Required columns not found: {PLAYER_ID_COLUMN}, {SLOT_REBATE_COLUMN}"
            ) from exc

        jobs = []
        for excel_row, row in enumerate(rows, start=2):
            player_value = row[player_idx] if player_idx < len(row) else None
            rebate_value = row[rebate_idx] if rebate_idx < len(row) else None
            if not numeric_nonzero(rebate_value):
                continue
            player_id = original_excel_text(player_value)
            amount = original_excel_text(rebate_value)
            if not player_id:
                raise ValueError(f"Missing Player ID at Excel row {excel_row}")
            jobs.append({"excel_row": excel_row, "player_id": player_id, "amount": amount})
        return jobs
    finally:
        wb.close()


# ------------------------------------------------------------ PM state

@dataclass
class Popup:
    hwnd: int
    title: str
    class_name: str
    kind: str

    def label(self):
        return f"{self.kind} '{self.title}' #{self.hwnd}"


@dataclass
class PMState:
    title: str
    main_enabled: bool
    popups: list

    def kinds(self):
        return sorted(p.kind for p in self.popups)

    def get(self, kind):
        for popup in self.popups:
            if popup.kind == kind:
                return popup
        return None

    def summary(self):
        popups = ", ".join(p.label() for p in self.popups) or "none"
        return f"title='{self.title}' main_enabled={self.main_enabled} popups=[{popups}]"


def classify(title, class_name):
    title = (title or "").strip()
    cls = (class_name or "").casefold()
    if LOGIN_TITLE_RE.match(title):
        return LOGIN
    if cls == "#32770":
        return DIALOG_TITLES.get(title.casefold(), UNKNOWN)
    if "popupbar" in cls or cls == "#32768":
        return MENU
    if cls == "combolbox":
        return DROPDOWN
    return UNKNOWN


def ids_in_title(title):
    return [x.lstrip("0") or "0" for x in re.findall(r"\((\d+)\)", title or "")]


def title_has_player(title, player_id):
    expected = str(player_id).lstrip("0") or "0"
    return expected in ids_in_title(title)


def normalized_number_text(value):
    value = str(value).strip().replace(",", "")
    try:
        number = float(value)
        return str(int(number)) if number.is_integer() else format(number, "g")
    except (TypeError, ValueError):
        return value


def escape_keys(text):
    return "".join("{%s}" % ch if ch in "{}[]()+^%~" else ch for ch in str(text))


# ------------------------------------------------ Win32 access to PM

class Win32PM:
    """Everything that touches the PM process lives here."""

    def __init__(self):
        if win32gui is None:
            raise SystemExit("This runner needs Windows with pywin32 and pywinauto installed.")
        self.pid = None
        self.main = None
        self._user32 = ctypes.windll.user32
        self._user32.GetAncestor.restype = ctypes.c_void_p
        self._user32.GetAncestor.argtypes = [ctypes.c_void_p, ctypes.c_uint]

    # -- discovery

    def connect(self):
        found = []

        def callback(hwnd, _):
            try:
                if (win32gui.IsWindowVisible(hwnd)
                        and win32gui.GetClassName(hwnd) in MAIN_CLASSES
                        and MAIN_TITLE_RE.match(win32gui.GetWindowText(hwnd))):
                    found.append(hwnd)
            except win32gui.error:
                pass
            return True

        win32gui.EnumWindows(callback, None)
        if not found:
            raise StepError("connect", "Patron Management main window not found. Open PM first.")
        if len(found) > 1:
            raise StepError("connect", "More than one Patron Management window is open.")
        self.main = found[0]
        self.pid = win32process.GetWindowThreadProcessId(self.main)[1]
        log(f"Connected to PM: main HWND={self.main}, PID={self.pid}")

    def main_title(self):
        return self.window_title(self.main)

    def main_enabled(self):
        return bool(win32gui.IsWindowEnabled(self.main))

    def popups(self):
        found = []

        def callback(hwnd, _):
            try:
                if hwnd == self.main or not win32gui.IsWindowVisible(hwnd):
                    return True
                if win32process.GetWindowThreadProcessId(hwnd)[1] != self.pid:
                    return True
                left, top, right, bottom = win32gui.GetWindowRect(hwnd)
                if right - left < MIN_POPUP_SIZE or bottom - top < MIN_POPUP_SIZE:
                    return True
                cls = win32gui.GetClassName(hwnd)
                if IGNORED_CLASS_RE.search(cls):
                    return True
                title = win32gui.GetWindowText(hwnd)
                found.append(Popup(hwnd, title, cls, classify(title, cls)))
            except win32gui.error:
                pass
            return True

        win32gui.EnumWindows(callback, None)
        return found

    # -- windows and controls

    def exists(self, hwnd):
        return bool(hwnd) and bool(win32gui.IsWindow(hwnd))

    def visible(self, hwnd):
        return self.exists(hwnd) and bool(win32gui.IsWindowVisible(hwnd))

    def enabled(self, hwnd):
        return self.visible(hwnd) and bool(win32gui.IsWindowEnabled(hwnd))

    def window_title(self, hwnd):
        try:
            return win32gui.GetWindowText(hwnd)
        except win32gui.error:
            return ""

    def child(self, parent, control_id):
        found = []

        def callback(hwnd, _):
            try:
                if win32gui.GetDlgCtrlID(hwnd) == control_id:
                    found.append(hwnd)
            except win32gui.error:
                pass
            return True

        try:
            win32gui.EnumChildWindows(parent, callback, None)
        except win32gui.error:
            pass
        visible = [h for h in found if win32gui.IsWindowVisible(h)]
        return (visible or found or [None])[0]

    def text(self, hwnd):
        try:
            return handleprops.text(hwnd) or ""
        except Exception:
            return ""

    def checked(self, hwnd):
        try:
            _, value = win32gui.SendMessageTimeout(
                hwnd, win32con.BM_GETCHECK, 0, 0, win32con.SMTO_ABORTIFHUNG, 2000)
            return value == win32con.BST_CHECKED
        except win32gui.error:
            return False

    def _top(self, hwnd):
        return int(self._user32.GetAncestor(hwnd, 2) or hwnd)

    def focus(self, hwnd):
        try:
            HwndWrapper(hwnd).set_focus()
        except Exception:
            pass

    def click(self, hwnd):
        """Physical click in the middle of a native control (no UIA Invoke)."""
        self.focus(self._top(hwnd))
        HwndWrapper(hwnd).click_input()
        time.sleep(0.2)

    def set_text(self, hwnd, value):
        HwndWrapper(hwnd).set_edit_text(str(value))

    def type_text(self, hwnd, value):
        self.click(hwnd)
        send_keys("{HOME}+{END}{DEL}")
        send_keys(escape_keys(value), with_spaces=True)

    def combo_items(self, hwnd):
        return list(HwndWrapper(hwnd).item_texts())

    def combo_select(self, hwnd, text):
        combo = HwndWrapper(hwnd)
        combo.select(combo.item_texts().index(text))   # exact match, never fuzzy

    def combo_selected(self, hwnd):
        combo = HwndWrapper(hwnd)
        index = combo.selected_index()
        items = combo.item_texts()
        return items[index] if 0 <= index < len(items) else ""

    def describe(self, hwnd):
        texts = []

        def callback(child, _):
            try:
                cls = win32gui.GetClassName(child)
                if cls in ("Static", "Button") and win32gui.IsWindowVisible(child):
                    text = self.text(child).strip()
                    if text:
                        texts.append(text)
            except win32gui.error:
                pass
            return True

        try:
            win32gui.EnumChildWindows(hwnd, callback, None)
        except win32gui.error:
            pass
        cls = win32gui.GetClassName(hwnd) if self.exists(hwnd) else "?"
        return f"'{self.window_title(hwnd)}' class={cls} hwnd={hwnd} texts={texts[:12]}"

    # -- UIA-only parts of PM

    def _uia(self, hwnd):
        return Desktop(backend="uia").window(handle=hwnd).wrapper_object()

    def click_ribbon_find(self):
        self.focus(self.main)
        for button in self._uia(self.main).descendants(control_type="Button"):
            if button.window_text().strip() == FIND_RIBBON_BUTTON:
                button.click_input()
                return True
        return False

    def send_find_shortcut(self):
        self.focus(self.main)
        send_keys(FIND_SHORTCUT)

    def click_named_item(self, popup_hwnd, name, control_type):
        for item in self._uia(popup_hwnd).descendants(control_type=control_type):
            if item.window_text().strip() == name:
                item.click_input()
                return True
        return False

    def click_menu_item(self, popup_hwnd, name):
        if self.click_named_item(popup_hwnd, name, "MenuItem"):
            return True
        # Fallback: recorded position of "Redeem Coupon..." in the 244x380 Options menu.
        left, top, right, bottom = win32gui.GetWindowRect(popup_hwnd)
        if name == REDEEM_MENU_ITEM and (right - left, bottom - top) == (244, 380):
            mouse.click(button="left", coords=(left + 122, top + 233))
            return True
        return False

    def list_items(self, popup_hwnd):
        try:
            return [clean for clean in (i.window_text().strip() for i in
                    self._uia(popup_hwnd).descendants(control_type="ListItem")) if clean]
        except Exception:
            return []

    def mdi_active(self):
        client = win32gui.FindWindowEx(self.main, 0, "MDIClient", None)
        if not client:
            return 0
        try:
            _, active = win32gui.SendMessageTimeout(
                client, WM_MDIGETACTIVE, 0, 0, win32con.SMTO_ABORTIFHUNG, 2000)
            return int(active or 0)
        except win32gui.error:
            return 0

    def close_tab(self, hwnd):
        win32gui.PostMessage(hwnd, win32con.WM_SYSCOMMAND, win32con.SC_CLOSE, 0)

    def fill_login(self, hwnd, username, password):
        self.focus(hwnd)
        window = self._uia(hwnd)
        edits = [e for e in window.descendants(control_type="Edit") if e.is_visible() and e.is_enabled()]

        def is_password(edit):
            try:
                return bool(edit.element_info.element.CurrentIsPassword)
            except Exception:
                return False

        passwords = [e for e in edits if is_password(e)]
        others = [e for e in edits if not is_password(e)]
        if passwords and others:
            user_ctl, password_ctl = others[0], passwords[0]
        elif len(edits) >= 2:
            user_ctl, password_ctl = edits[0], edits[1]
        else:
            raise StepError("0_login", "Could not identify separate Username and Password fields.")
        for control, value in ((user_ctl, username), (password_ctl, password)):
            control.click_input()
            send_keys("{HOME}+{END}{DEL}")
            send_keys(escape_keys(value), with_spaces=True)
        typed_user = ""
        try:
            typed_user = user_ctl.get_value() or ""
        except Exception:
            typed_user = username
        if typed_user.strip() != username.strip():
            raise StepError("0_login", "Username field does not show the configured username.")
        buttons = [b for b in window.descendants(control_type="Button")
                   if b.window_text().strip().casefold() in LOGIN_BUTTON_NAMES and b.is_enabled()]
        if buttons:
            buttons[0].click_input()
        else:
            send_keys("{ENTER}")


# ------------------------------------------------------ gates (pre/post)

def read_state(pm):
    return PMState(pm.main_title(), pm.main_enabled(), pm.popups())


def state_problem(state, popups, main_enabled, player_id):
    expected = sorted(popups)
    if state.kinds() != expected:
        return f"expected popups {expected or 'none'}, found {state.kinds() or 'none'}"
    if main_enabled is not None and state.main_enabled != main_enabled:
        return f"expected main window {'enabled' if main_enabled else 'blocked'}"
    if player_id is not None and not title_has_player(state.title, player_id):
        return f"main title does not show player {player_id}: '{state.title}'"
    return None


def raise_on_unknown(pm, step, state, while_doing):
    unknown = [p for p in state.popups if p.kind == UNKNOWN]
    if unknown:
        raise StepError(step, f"unexpected popup while {while_doing}: {pm.describe(unknown[0].hwnd)}")


def check_state(pm, step, popups=(), main_enabled=None, player_id=None):
    """Pre-check: PM must already be in this state, otherwise stop."""
    state = read_state(pm)
    raise_on_unknown(pm, step, state, "checking the state before this step")
    problem = state_problem(state, popups, main_enabled, player_id)
    if problem:
        raise StepError(step, f"pre-check failed: {problem}. State: {state.summary()}")
    log(f"  [{step}] pre-check OK   {state.summary()}")
    return state


def wait_state(pm, step, popups=(), main_enabled=None, player_id=None,
               timeout=T_VERIFY, quiet=0.0, what="the expected state", fail=True):
    """Post-check: wait until PM reaches this state (stable for `quiet` seconds)."""
    deadline = time.time() + timeout
    stable_since = None
    problem = "not checked yet"
    while True:
        state = read_state(pm)
        raise_on_unknown(pm, step, state, f"waiting for {what}")
        problem = state_problem(state, popups, main_enabled, player_id)
        if problem is None:
            stable_since = stable_since or time.time()
            if time.time() - stable_since >= quiet:
                log(f"  [{step}] post-check OK  {state.summary()}")
                return state
        else:
            stable_since = None
        if time.time() > deadline:
            if not fail:
                return None
            raise StepError(step, f"timed out after {timeout}s waiting for {what}: {problem}. "
                                  f"State: {state.summary()}")
        time.sleep(POLL)


def wait_until(predicate, timeout, interval=POLL):
    deadline = time.time() + timeout
    while True:
        try:
            if predicate():
                return True
        except Exception:
            pass
        if time.time() > deadline:
            return False
        time.sleep(interval)


def wait_popup_closed(pm, step, popup, timeout=T_POPUP_CLOSE):
    deadline = time.time() + timeout
    while pm.visible(popup.hwnd):
        state = read_state(pm)
        raise_on_unknown(pm, step, state, f"waiting for {popup.kind} to close")
        if time.time() > deadline:
            raise StepError(step, f"{popup.label()} did not close within {timeout}s. State: {state.summary()}")
        time.sleep(POLL)
    log(f"  [{step}] {popup.kind} closed")


def control(pm, step, popup, control_id, what, need_enabled=True, timeout=T_VERIFY):
    found = {}

    def ready():
        hwnd = pm.child(popup.hwnd, control_id)
        found["hwnd"] = hwnd
        return hwnd and (pm.enabled(hwnd) if need_enabled else pm.exists(hwnd))

    if not wait_until(ready, timeout):
        state = "missing" if not found.get("hwnd") else "disabled"
        raise StepError(step, f"{what} (id={control_id}) in {popup.label()} is {state}.")
    return found["hwnd"]


def click(pm, step, popup, control_id, what, allow_coupon_ok=False):
    allowed = COUPON_CLICKABLE_IDS | ({ID_OK} if allow_coupon_ok else set())
    if popup.kind == COUPON and control_id not in allowed:
        raise StepError(step, f"Safety stop: refusing to click control id={control_id} in Coupon Redemption.")
    hwnd = control(pm, step, popup, control_id, what)
    log(f"  [{step}] click {what} (id={control_id})")
    pm.click(hwnd)
    return hwnd


def set_and_verify(pm, step, hwnd, value, what, numeric=False):
    norm = normalized_number_text if numeric else (lambda v: str(v).strip())
    expected = norm(value)

    def matches():
        return norm(pm.text(hwnd)) == expected

    pm.set_text(hwnd, value)
    if not wait_until(matches, 1.5):
        log(f"  [{step}] {what}: direct set did not stick, typing instead")
        pm.type_text(hwnd, value)
        if not wait_until(matches, T_VERIFY):
            raise StepError(step, f"{what} shows '{pm.text(hwnd)}', expected '{expected}'.")
    log(f"  [{step}] checkpoint: {what} = {expected}")


# ----------------------------------------------------------------- steps

def ensure_logged_in(pm, username, password):
    step = "0_login"
    state = read_state(pm)
    if state.get(LOGIN) is None and state.title.startswith("Patron Management - "):
        log("PM is already logged in.")
        return
    if state.get(LOGIN) is None:
        log("Waiting for the PM login window...")
        state = wait_state(pm, step, popups=(LOGIN,), timeout=30, what="the login window")
    if (not username or not password or username == "ENTER_USERNAME_HERE"
            or password == "ENTER_PASSWORD_HERE"):
        raise StepError(step, "Fill in pm_credentials.py first.")
    login_window = check_state(pm, step, popups=(LOGIN,)).get(LOGIN)
    log(f"  [{step}] typing credentials into {login_window.label()}")
    pm.fill_login(login_window.hwnd, username, password)
    deadline = time.time() + T_LOGIN
    while True:
        state = read_state(pm)
        raise_on_unknown(pm, step, state, "logging in")
        if (state.get(LOGIN) is None and not state.popups and state.main_enabled
                and state.title.startswith("Patron Management - ")):
            break
        if time.time() > deadline:
            raise StepError(step, f"login did not finish within {T_LOGIN}s. State: {state.summary()}")
        time.sleep(POLL)
    wait_state(pm, step, popups=(), main_enabled=True, quiet=QUIET_SECONDS, timeout=10,
               what="PM idle after login")
    log("Login completed.")


def open_find_player(pm, step):
    check_state(pm, step, popups=(), main_enabled=True)
    clicked = False
    try:
        log(f"  [{step}] click Ribbon '{FIND_RIBBON_BUTTON}'")
        clicked = pm.click_ribbon_find()
    except Exception as exc:
        log(f"  [{step}] Ribbon click failed: {exc}")
    if clicked:
        state = wait_state(pm, step, popups=(FIND,), main_enabled=False, timeout=T_FIND_OPEN,
                           what="Find a Player", fail=False)
        if state:
            return state
    state = read_state(pm)
    if state.popups or not state.main_enabled:
        # Something is already opening; do not send a second command.
        return wait_state(pm, step, popups=(FIND,), main_enabled=False, timeout=T_FIND_OPEN,
                          what="Find a Player")
    log(f"  [{step}] fallback: shortcut {FIND_SHORTCUT}")
    pm.send_find_shortcut()
    return wait_state(pm, step, popups=(FIND,), main_enabled=False, timeout=T_FIND_OPEN,
                      what="Find a Player")


def wait_profile_title(pm, step, player_id, previous_title, previous_tab):
    """Wait for the new profile; the old title stays for ~4-5s and is not an error.

    When this player was already on screen the title cannot show the reload, so
    wait for a new profile tab, or SAME_PLAYER_WAIT_SECONDS if PM reuses the tab.
    """
    started = time.time()
    deadline = started + T_PROFILE_LOAD
    same_player = title_has_player(previous_title, player_id)
    while True:
        state = read_state(pm)
        raise_on_unknown(pm, step, state, "waiting for the profile to load")
        others = [p for p in state.popups if p.kind not in (SYSMSG, COMMENT)]
        if others:
            raise StepError(step, f"unexpected popup while loading the profile: {others[0].label()}")
        if title_has_player(state.title, player_id):
            tab = pm.mdi_active()
            if not same_player or (tab and tab != previous_tab):
                log(f"  [{step}] post-check OK  profile loaded: '{state.title}'")
                return state
            if time.time() - started >= SAME_PLAYER_WAIT_SECONDS:
                log(f"  [{step}] post-check OK  same player was already open; "
                    f"waited {SAME_PLAYER_WAIT_SECONDS:.0f}s for the reload")
                return state
        if state.title != previous_title and ids_in_title(state.title):
            raise StepError(step, f"Wrong profile opened for {player_id}: '{state.title}'")
        if time.time() > deadline:
            raise StepError(step, f"profile {player_id} did not load within {T_PROFILE_LOAD}s. "
                                  f"State: {state.summary()}")
        time.sleep(POLL)


def handle_system_messages(pm, step, popup):
    items = pm.list_items(popup.hwnd)
    log(f"  [{step}] System Messages: {' | '.join(items) or '(no readable items)'}")
    if STOP_CODES_ITEM in items:
        log(f"  [{step}] NOTE: player has '{STOP_CODES_ITEM}' (logged only, continuing)")
    click(pm, step, popup, SYSMSG_CLOSE, "System Messages Close")
    wait_popup_closed(pm, step, popup)
    return items


def comment_signature(pm, popup, next_hwnd, close_hwnd):
    header = pm.child(popup.hwnd, COMMENT_HEADER)
    previous = pm.child(popup.hwnd, COMMENT_PREVIOUS)
    return (pm.text(header) if header else "", pm.enabled(next_hwnd),
            pm.enabled(previous) if previous else None, pm.enabled(close_hwnd))


def handle_player_comment(pm, step, popup):
    """Close is disabled until the last comment has been shown: page with Next first."""
    next_hwnd = control(pm, step, popup, COMMENT_NEXT, "Player Comment Next", need_enabled=False)
    close_hwnd = control(pm, step, popup, COMMENT_CLOSE, "Player Comment Close", need_enabled=False)
    page = 1
    while True:
        if pm.enabled(close_hwnd):
            log(f"  [{step}] Player Comment: all {page} page(s) shown, Close is enabled")
            click(pm, step, popup, COMMENT_CLOSE, "Player Comment Close")
            wait_popup_closed(pm, step, popup)
            return page
        if page >= MAX_COMMENT_PAGES:
            raise StepError(step, f"Player Comment still not closable after {page} pages.")
        if pm.enabled(next_hwnd):
            before = comment_signature(pm, popup, next_hwnd, close_hwnd)
            for attempt in (1, 2):
                log(f"  [{step}] Player Comment: click Next (page {page} -> {page + 1})")
                pm.click(next_hwnd)
                if wait_until(lambda: comment_signature(pm, popup, next_hwnd, close_hwnd) != before,
                              T_COMMENT_PAGE):
                    break
                if attempt == 2:
                    raise StepError(step, "Player Comment did not react to Next.")
            page += 1
            continue
        # Next can be hidden for ~0.1s while PM redraws it.
        if not wait_until(lambda: pm.enabled(close_hwnd) or pm.enabled(next_hwnd), 2.0):
            raise StepError(step, "Player Comment: Next and Close are both disabled.")


def settle_profile(pm, step, player_id, quiet):
    """Handle System Messages / Player Comment until PM stays idle for `quiet` seconds."""
    deadline = time.time() + T_PROFILE_POPUPS
    quiet_since = None
    messages, pages = [], 0
    while True:
        state = read_state(pm)
        raise_on_unknown(pm, step, state, "handling profile popups")
        if not title_has_player(state.title, player_id):
            raise StepError(step, f"profile title changed unexpectedly: '{state.title}'")
        others = [p for p in state.popups if p.kind not in (SYSMSG, COMMENT)]
        if others:
            raise StepError(step, f"unexpected popup on the profile: {others[0].label()}")
        if state.get(SYSMSG):
            messages.extend(handle_system_messages(pm, step, state.get(SYSMSG)))
            quiet_since = None
        elif state.get(COMMENT):
            pages += handle_player_comment(pm, step, state.get(COMMENT))
            quiet_since = None
        elif not state.main_enabled:
            quiet_since = None      # a modal popup is being created (shows 0.03-0.25s later)
        else:
            quiet_since = quiet_since or time.time()
            if time.time() - quiet_since >= quiet:
                log(f"  [{step}] post-check OK  profile idle for {quiet:.1f}s  {state.summary()}")
                return messages, pages
        if time.time() > deadline:
            raise StepError(step, f"profile did not become idle within {T_PROFILE_POPUPS}s. "
                                  f"State: {state.summary()}")
        time.sleep(POLL)


def active_profile_tab(pm, step, player_id):
    tab = pm.mdi_active()
    title = pm.window_title(tab) if tab else ""
    if not tab or not title_has_player(title, player_id):
        raise StepError(step, f"active profile tab is not player {player_id}: '{title}'")
    return tab


def select_competitor(pm, step, coupon, combo_hwnd):
    items = pm.combo_items(combo_hwnd)
    if COMPETITOR_NAME not in items:
        raise StepError(step, f"'{COMPETITOR_NAME}' is not in the Competitor list ({len(items)} items).")
    log(f"  [{step}] select '{COMPETITOR_NAME}' in Competitor list")
    try:
        pm.combo_select(combo_hwnd, COMPETITOR_NAME)
    except Exception as exc:
        log(f"  [{step}] direct select failed: {exc}")
    if wait_until(lambda: pm.combo_selected(combo_hwnd) == COMPETITOR_NAME, T_VERIFY):
        return
    log(f"  [{step}] fallback: open the list and click the item")
    pm.click(combo_hwnd)
    state = wait_state(pm, step, popups=(COUPON, DROPDOWN), timeout=T_VERIFY, what="the Competitor list")
    if not pm.click_named_item(state.get(DROPDOWN).hwnd, COMPETITOR_NAME, "ListItem"):
        raise StepError(step, f"'{COMPETITOR_NAME}' not found in the open Competitor list.")
    wait_state(pm, step, popups=(COUPON,), timeout=T_VERIFY, what="the Competitor list to close")
    if not wait_until(lambda: pm.combo_selected(combo_hwnd) == COMPETITOR_NAME, T_VERIFY):
        raise StepError(step, f"Competitor shows '{pm.combo_selected(combo_hwnd)}', "
                              f"expected '{COMPETITOR_NAME}'.")


def begin_step(progress, name):
    """Pause STEP_DELAY_SECONDS before every step, then record the step name."""
    time.sleep(STEP_DELAY_SECONDS)
    progress["step"] = name
    return name


def process_player(pm, job, progress, allow_ok=False):
    player_id, amount = job["player_id"], job["amount"]

    step = begin_step(progress, "1_precheck_idle")
    state = check_state(pm, step, popups=(), main_enabled=True)
    if not state.title.startswith("Patron Management - "):
        raise StepError(step, f"PM is not on a logged-in page: '{state.title}'")
    previous_title = state.title
    previous_tab = pm.mdi_active()

    step = begin_step(progress, "2_open_find_player")
    open_find_player(pm, step)

    step = begin_step(progress, "3_enter_player_id")
    find = check_state(pm, step, popups=(FIND,), main_enabled=False).get(FIND)
    field = control(pm, step, find, FIND_PLAYER_ID_EDIT, "Player ID field")
    set_and_verify(pm, step, field, player_id, "Player ID")

    step = begin_step(progress, "4_confirm_find")
    find = check_state(pm, step, popups=(FIND,), main_enabled=False).get(FIND)
    if pm.text(field).strip() != player_id:
        raise StepError(step, f"Player ID field changed to '{pm.text(field)}' before OK.")
    click(pm, step, find, ID_OK, "Find a Player OK")
    wait_popup_closed(pm, step, find, T_FIND_CLOSE)

    step = begin_step(progress, "5_wait_profile_loaded")
    wait_profile_title(pm, step, player_id, previous_title, previous_tab)

    step = begin_step(progress, "6_profile_popups")
    messages, pages = settle_profile(pm, step, player_id, QUIET_SECONDS)
    tab = active_profile_tab(pm, step, player_id)
    progress["tab"] = tab
    progress["system_messages"] = " | ".join(messages)
    progress["stop_codes"] = "YES" if STOP_CODES_ITEM in messages else "NO"
    progress["comment_pages"] = pages

    step = begin_step(progress, "7_open_options_menu")
    check_state(pm, step, popups=(), main_enabled=True, player_id=player_id)
    tab = active_profile_tab(pm, step, player_id)
    options = pm.child(tab, OPTIONS_BUTTON)
    if not options or not pm.enabled(options):
        raise StepError(step, "Options... button not found or disabled on the active profile tab.")
    log(f"  [{step}] click Options... (id={OPTIONS_BUTTON})")
    pm.click(options)
    wait_state(pm, step, popups=(MENU,), main_enabled=True, player_id=player_id,
               timeout=T_MENU_OPEN, what="the Options menu")

    step = begin_step(progress, "8_click_redeem_coupon")
    menu = check_state(pm, step, popups=(MENU,), main_enabled=True, player_id=player_id).get(MENU)
    log(f"  [{step}] click menu item '{REDEEM_MENU_ITEM}'")
    if not pm.click_menu_item(menu.hwnd, REDEEM_MENU_ITEM):
        raise StepError(step, f"Menu item '{REDEEM_MENU_ITEM}' not found.")
    wait_state(pm, step, popups=(COUPON,), main_enabled=False, player_id=player_id,
               timeout=T_COUPON_OPEN, what="Coupon Redemption")

    step = begin_step(progress, "9_select_competitor_coupon")
    coupon = check_state(pm, step, popups=(COUPON,), main_enabled=False, player_id=player_id).get(COUPON)
    radio = control(pm, step, coupon, COUPON_COMPETITOR_RADIO, "Competitor Coupon")
    combo = control(pm, step, coupon, COUPON_COMPETITOR_COMBO, "Competitor list", need_enabled=False)
    amount_field = control(pm, step, coupon, COUPON_AMOUNT_EDIT, "Amount", need_enabled=False)
    log(f"  [{step}] initial: Competitor list enabled={pm.enabled(combo)}, Amount enabled={pm.enabled(amount_field)}")
    click(pm, step, coupon, COUPON_COMPETITOR_RADIO, "Competitor Coupon")
    if not wait_until(lambda: pm.checked(radio) and pm.enabled(combo) and pm.enabled(amount_field),
                      T_FIELDS_ENABLE):
        raise StepError(step, f"after Competitor Coupon: selected={pm.checked(radio)}, "
                              f"list enabled={pm.enabled(combo)}, Amount enabled={pm.enabled(amount_field)}")
    wait_state(pm, step, popups=(COUPON,), main_enabled=False, player_id=player_id,
               what="Coupon Redemption only")
    log(f"  [{step}] checkpoint: Competitor Coupon selected, list and Amount enabled")

    step = begin_step(progress, "10_select_competitor")
    check_state(pm, step, popups=(COUPON,), main_enabled=False, player_id=player_id)
    select_competitor(pm, step, coupon, combo)
    wait_state(pm, step, popups=(COUPON,), main_enabled=False, player_id=player_id,
               what="Coupon Redemption only")
    log(f"  [{step}] checkpoint: Competitor = {COMPETITOR_NAME}")

    step = begin_step(progress, "11_enter_amount")
    check_state(pm, step, popups=(COUPON,), main_enabled=False, player_id=player_id)
    set_and_verify(pm, step, amount_field, amount, "Amount", numeric=True)
    wait_state(pm, step, popups=(COUPON,), main_enabled=False, player_id=player_id,
               quiet=0.5, what="no popup after entering the amount")

    step = begin_step(progress, "12_ok_coupon" if allow_ok else "12_cancel_coupon")
    check_state(pm, step, popups=(COUPON,), main_enabled=False, player_id=player_id)
    if pm.combo_selected(combo) != COMPETITOR_NAME or \
            normalized_number_text(pm.text(amount_field)) != normalized_number_text(amount) or \
            not pm.checked(radio) or active_profile_tab(pm, step, player_id) != progress["tab"]:
        raise StepError(step, "Coupon fields or profile changed before the final click. Nothing was clicked.")
    log(f"  [{step}] final check OK: player {player_id}, {COMPETITOR_NAME}, amount {amount}")
    if not allow_ok:
        click(pm, step, coupon, ID_CANCEL, "Coupon Redemption Cancel")
        wait_popup_closed(pm, step, coupon)
        wait_state(pm, step, popups=(), main_enabled=True, player_id=player_id,
                   timeout=T_POPUP_CLOSE, quiet=0.5, what="PM idle after Cancel")
        return
    # Written before the click: if anything goes wrong afterwards this player is
    # still treated as redeemed and is never clicked again by a rerun today.
    ledger_append(job, "OK_CLICKED")
    progress["ok_clicked"] = True
    click(pm, step, coupon, ID_OK, "Coupon Redemption OK", allow_coupon_ok=True)
    wait_popup_closed(pm, step, coupon, T_AFTER_OK)
    wait_state(pm, step, popups=(), main_enabled=True, player_id=player_id,
               timeout=T_AFTER_OK, quiet=1.0, what="PM idle after OK")
    ledger_append(job, "REDEEMED")


def close_profile_tabs(pm, tabs):
    step = "13_close_profile_tabs"
    log(f"Closing {len(tabs)} profile tab(s) opened by this run.")
    closed = 0
    for tab, player_id in tabs:
        if not pm.exists(tab):
            continue
        title = pm.window_title(tab)
        if not title_has_player(title, player_id):
            log(f"  [{step}] skip tab #{tab}: title '{title}' is not player {player_id}")
            continue
        check_state(pm, step, popups=(), main_enabled=True)
        log(f"  [{step}] close tab '{title}'")
        pm.close_tab(tab)
        if not wait_until(lambda: not pm.exists(tab) or not pm.visible(tab), T_TAB_CLOSE):
            raise StepError(step, f"tab '{title}' did not close. State: {read_state(pm).summary()}")
        wait_state(pm, step, popups=(), main_enabled=True, timeout=T_TAB_CLOSE, quiet=0.5,
                   what="PM idle after closing the tab")
        closed += 1
    log(f"Profile tabs closed: {closed}")


# ------------------------------------------------------------------ run

LEDGER_FIELDS = ["timestamp", "date", "player_id", "amount", "competitor", "status", "excel_row"]


def ledger_players_today():
    """Player IDs with an OK click recorded today (OK_CLICKED or REDEEMED)."""
    if not LEDGER_FILE.exists():
        return set()
    today = datetime.now().date().isoformat()
    with LEDGER_FILE.open(newline="", encoding="utf-8-sig") as file:
        return {row["player_id"] for row in csv.DictReader(file) if row.get("date") == today}


def ledger_append(job, status):
    new_file = not LEDGER_FILE.exists()
    with LEDGER_FILE.open("a", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=LEDGER_FIELDS)
        if new_file:
            writer.writeheader()
        now = datetime.now()
        writer.writerow({"timestamp": now.isoformat(timespec="seconds"), "date": now.date().isoformat(),
                         "player_id": job["player_id"], "amount": job["amount"],
                         "competitor": COMPETITOR_NAME, "status": status, "excel_row": job["excel_row"]})


def ask_allow_ok(jobs):
    total = sum(float(job["amount"]) for job in jobs)
    print("\n" + "=" * 70)
    print(f"{len(jobs)} player(s), total SLOT REBATE 5% = {normalized_number_text(total)}")
    for job in jobs:
        print(f"   row {job['excel_row']:>4}  player {job['player_id']:>8}  amount {job['amount']}")
    print("=" * 70)
    answer = input("Allow clicking OK in Coupon Redemption (REAL redemption)?\n"
                   "Type YES to click OK, anything else = Cancel only (test run): ")
    return answer.strip() == "YES"


CSV_FIELDS = [
    "timestamp", "excel_row", "player_id", "slot_rebate_5_percent", "test_mode", "status",
    "failed_step", "message", "stop_codes", "system_messages", "comment_pages", "duration_s",
]


def write_log(rows):
    with LOG_FILE.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=CSV_FIELDS)
        writer.writeheader()
        writer.writerows(rows)


def run(pm, jobs, username, password, allow_ok=False):
    """Process every job in order; stop at the first error and leave PM untouched."""
    log(f"Mode: {'OK - REAL REDEMPTION' if allow_ok else 'Cancel only (test run)'}")
    done_today = ledger_players_today() if allow_ok else set()
    results = []
    tabs = []
    progress = {"step": "0_login"}
    try:
        ensure_logged_in(pm, username, password)
    except Exception as exc:
        log(f"STOPPED at login: {exc}")
        write_log(results)
        return 1

    for index, job in enumerate(jobs, start=1):
        started = time.time()
        progress = {"step": "1_precheck_idle"}
        record = {
            "timestamp": datetime.now().isoformat(timespec="seconds"),
            "excel_row": job["excel_row"],
            "player_id": job["player_id"],
            "slot_rebate_5_percent": job["amount"],
            "test_mode": not allow_ok,
            "status": "", "failed_step": "", "message": "",
            "stop_codes": "", "system_messages": "", "comment_pages": "", "duration_s": "",
        }
        log(f"[{index}/{len(jobs)}] Player ID {job['player_id']}, Amount {job['amount']} "
            f"(Excel row {job['excel_row']})")
        if job["player_id"] in done_today:
            record.update({"status": "SKIPPED_ALREADY_REDEEMED", "duration_s": 0,
                           "message": f"OK was already clicked today (see {LEDGER_FILE.name})."})
            results.append(record)
            write_log(results)
            log(f"Player {job['player_id']}: skipped, already redeemed today")
            continue
        try:
            process_player(pm, job, progress, allow_ok)
            if allow_ok:
                record["status"] = "REDEEMED"
                record["message"] = "Amount entered from Excel; Coupon Redemption confirmed with OK."
            else:
                record["status"] = "TEST_CANCELLED"
                record["message"] = "Amount entered from Excel; Coupon Redemption cancelled."
        except Exception as exc:
            failed_step = getattr(exc, "step", progress["step"])
            try:
                screen = read_state(pm).summary()
            except Exception:
                screen = "unavailable"
            record.update({
                "status": "ERROR_AFTER_OK" if progress.get("ok_clicked") else "ERROR",
                "failed_step": failed_step,
                "message": f"{type(exc).__name__}: {exc} | screen: {screen}",
            })
        record["stop_codes"] = progress.get("stop_codes", "")
        record["system_messages"] = progress.get("system_messages", "")
        record["comment_pages"] = progress.get("comment_pages", "")
        record["duration_s"] = round(time.time() - started, 1)
        results.append(record)
        write_log(results)
        if progress.get("tab"):
            tabs.append((progress["tab"], job["player_id"]))
        if record["status"].startswith("ERROR"):
            log(f"Player {job['player_id']}: ERROR at {record['failed_step']}: {record['message']}")
            log("STOPPED. PM is left exactly as it is for inspection; no cleanup was done.")
            log(f"Log file: {LOG_FILE}")
            return 1
        log(f"Player {job['player_id']}: {record['status']} in {record['duration_s']}s")

    if CLOSE_TABS_AT_END:
        try:
            close_profile_tabs(pm, tabs)
        except Exception as exc:
            log(f"STOPPED while closing tabs at {getattr(exc, 'step', 'close_tabs')}: {exc}")
            return 1
    counts = {}
    for row in results:
        counts[row["status"]] = counts.get(row["status"], 0) + 1
    log(f"Done. {counts}")
    log(f"Log file: {LOG_FILE}")
    return 0


def main():
    try:
        from pm_credentials import PM_USERNAME, PM_PASSWORD
    except ImportError as exc:
        raise SystemExit("Missing pm_credentials.py in the same folder.") from exc

    jobs = read_jobs()
    log(f"Runner {RUNNER_VERSION}")
    if not jobs:
        log("No non-zero SLOT REBATE 5% rows found.")
        write_log([])
        return 0
    log(f"Loaded {len(jobs)} Slot Rebate job(s) from Excel: " + ", ".join(x["player_id"] for x in jobs))
    allow_ok = ask_allow_ok(jobs)
    pm = Win32PM()
    pm.connect()
    return run(pm, jobs, PM_USERNAME, PM_PASSWORD, allow_ok)


if __name__ == "__main__":
    try:
        sys.exit(main())
    except SystemExit:
        raise
    except Exception as exc:
        print(f"[FATAL] {type(exc).__name__}: {exc}")
        sys.exit(1)
